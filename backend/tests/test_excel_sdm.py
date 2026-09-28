"""Deterministic Excel SDM content; not model, engine or release acceptance."""
from copy import deepcopy
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from app.sdm_specification import build_sdm_candidate
from app.sdm_renderer import render_sdm_xlsx
from app.sdm_xlsx_semantics import expected_sdm_cells, validate_sdm_xlsx
from test_excel_specification import excel_design
from app.excel_contract_binding import validated_excel_contract


@pytest.mark.parametrize('blank_rows', ['SKIP', 'PRESERVE'])
def test_excel_sdm_binds_confirmed_policy_without_upload_data(blank_rows):
    spec,run,naming=excel_design()
    config=run['input_snapshot']['source_config']
    config['excel_input_contract_v1']['blank_rows']=blank_rows
    bound=validated_excel_contract(config)
    spec['excel_source']=bound['reference']
    before=deepcopy((spec,run,naming))
    candidate=build_sdm_candidate(spec,run,naming)
    document=candidate['document']
    assert document['version']==4 and document['source_format']=='XLSX'
    assert document['excel_source']==bound['reference']
    assert document['excel_input_contract']==bound['contract']
    assert candidate==build_sdm_candidate(spec,run,naming)
    assert (spec,run,naming)==before
    assert config['sources'][0]['upload_id'] not in json.dumps(document)
    assert not candidate['qa_passed'] and not candidate['release_ready']
    rendered=render_sdm_xlsx(spec,run,naming)
    assert rendered['content']==render_sdm_xlsx(spec,run,naming)['content']
    assert rendered['evidence']['semantic_equality']=='CANDIDATE_LAYOUT_MATCHED'
    book=load_workbook(BytesIO(rendered['content']))
    rules=dict(book['規則及版本'].iter_rows(min_row=6,values_only=True))
    assert rules['Excel 工作表']=='明細'
    assert rules['Excel 標頭列'].startswith('2（')
    assert rules['空白列政策'].startswith(blank_rows)
    assert rules['Excel 讀取契約 SHA-256']==bound['reference']['contract_checksum']
    assert 'FAIL' in rules['型別與錯誤']
    book.close()


@pytest.mark.parametrize('key',['content_checksum','profile_checksum','contract_checksum'])
def test_excel_sdm_rejects_changed_specification_reference(key):
    args=excel_design()
    args[0]['excel_source'][key]='0'*64
    with pytest.raises(ValueError,match='SDM_VALID_SPECIFICATION_REQUIRED'):
        render_sdm_xlsx(*args)


def test_excel_sdm_rejects_policy_change_and_document_from_old_policy():
    args=excel_design()
    old=render_sdm_xlsx(*args)['content']
    args[1]['input_snapshot']['source_config']['excel_input_contract_v1']['blank_rows']='PRESERVE'
    with pytest.raises(ValueError,match='SDM_VALID_SPECIFICATION_REQUIRED'):
        expected_sdm_cells(*args)
    args[0]['excel_source']=validated_excel_contract(args[1]['input_snapshot']['source_config'])['reference']
    with pytest.raises(ValueError,match='SDM_XLSX_SPECIFICATION_MISMATCH'):
        validate_sdm_xlsx(old,*args)
