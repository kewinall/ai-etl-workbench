from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import task_uploads
from app.excel_profile import inspect_workbook, selected_profile, verify_excel_source, create_excel_profile_router
from app.source_preflight import source_preflight
from app.source_profiler import profile_task


def workbook_bytes():
    book = Workbook()
    book.active.title = '說明'
    book.active.append(['不要自動選我'])
    sheet = book.create_sheet('明細')
    sheet.append(['標題列在下一列'])
    sheet.append(['客戶編號', '金額'])
    sheet.append(['001', '12.123456'])
    bad = book.create_sheet('重複欄')
    bad.append(['編號', '編號'])
    bad.append(['A', 'B'])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.fixture
def uploaded(tmp_path, monkeypatch):
    monkeypatch.setattr(task_uploads, 'UPLOAD_ROOT', tmp_path)
    return task_uploads.save_and_profile('synthetic.xlsx', workbook_bytes())


def test_upload_lists_sheets_without_selecting_first(uploaded):
    assert uploaded['worksheets'] == ['說明', '明細', '重複欄']
    assert uploaded['worksheet'] is None
    assert uploaded['fields'] == []
    assert uploaded['selection_required'] is True


def test_selection_reprofiles_and_binds_exact_columns(uploaded):
    selected = selected_profile(uploaded['upload_id'], uploaded['checksum'], uploaded['size'], '明細', 2)
    assert selected['fields'] == [{'name': '客戶編號', 'type': 'VARCHAR(32)'},
                                  {'name': '金額', 'type': 'DECIMAL(18,6)'}]
    source = {**uploaded, **selected, 'type': 'EXCEL', 'has_actual_data': True}
    assert verify_excel_source(source) == selected
    assert selected['full_file_validated'] is False
    profile = profile_task({'source': 'EXCEL', 'source_config': {'sources': [source]}})
    assert profile['sources'][0]['fields'] == selected['fields']
    evidence, issues = source_preflight({'source_config': {'sources': [source]}})
    assert issues == []
    assert evidence[0]['excel_selection_v1'] == selected['excel_selection_v1']
    assert evidence[0]['execution_authorized'] is False


@pytest.mark.parametrize('field,value', [('worksheet', '說明'), ('header_row', 1),
                                       ('checksum', '0' * 64), ('fields', []),
                                       ('excel_selection_v1', None)])
def test_changed_input_cannot_reuse_confirmation(uploaded, field, value):
    selected = selected_profile(uploaded['upload_id'], uploaded['checksum'], uploaded['size'], '明細', 2)
    source = {**uploaded, **selected, 'type': 'EXCEL', field: value}
    with pytest.raises(ValueError):
        verify_excel_source(source)


@pytest.mark.parametrize('sheet,header', [('明細', 1), ('重複欄', 1), ('不存在', 1), ('明細', 0), ('明細', True)])
def test_missing_duplicate_and_invalid_headers_rejected(sheet, header):
    with pytest.raises(ValueError):
        inspect_workbook(workbook_bytes(), sheet, header)


def test_endpoint_rejects_arbitrary_path_and_validates_selection(uploaded):
    app = FastAPI()
    app.include_router(create_excel_profile_router())
    client = TestClient(app)
    url = f"/api/task-sources/{uploaded['upload_id']}/excel-profile"
    data = {'checksum': uploaded['checksum'], 'size': uploaded['size'], 'worksheet': '明細', 'header_row': 2}
    assert client.post(url, json=data).status_code == 200
    assert client.post(url, json={**data, 'path': '/not-a-managed-upload.xlsx'}).status_code == 422
    assert client.post(url, json={**data, 'header_row': '2'}).status_code == 422
    assert client.post(url, json={**data, 'size': uploaded['size'] + 1}).status_code == 422


def test_preflight_uses_the_already_verified_byte_snapshot(uploaded, monkeypatch):
    from app import upload_integrity
    selected = selected_profile(uploaded['upload_id'], uploaded['checksum'], uploaded['size'], '明細', 2)
    source = {**uploaded, **selected, 'type': 'EXCEL'}
    # use the actual immutable upload, not a newly rendered XLSX ZIP timestamp
    from pathlib import Path
    content = Path(uploaded['path']).read_bytes()
    def unexpected_second_read(*args):
        raise AssertionError('must not reopen a verified byte snapshot')
    monkeypatch.setattr(upload_integrity, 'read_verified_upload', unexpected_second_read)
    assert verify_excel_source(source, content=content) == selected
    with pytest.raises(ValueError, match='UPLOAD_CONTENT_CHANGED'):
        verify_excel_source(source, content=content + b'changed')
