from io import BytesIO
from openpyxl import load_workbook
import pytest
from app.sdm_renderer import render_sdm_xlsx
from app.sdm_xlsx_semantics import validate_sdm_xlsx
from app.sa_contract import digest
from app import release_files
from test_etl_specification import design


def binding(spec):
    return dict(version=1,run_id=spec['run_id'],specification_checksum=digest(spec),qa_binding_checksum='a'*64,
                source_sdm_checksum='b'*64,portability_checksum='c'*64)


def test_delivery_document_preserves_mapping_and_links_evidence():
    spec,run,naming=design();bound=binding(spec)
    candidate=render_sdm_xlsx(spec,run,naming)['content']
    content=render_sdm_xlsx(spec,run,naming,release_binding=bound)['content']
    assert content!=candidate
    assert render_sdm_xlsx(spec,run,naming,release_binding=bound)['content']==content
    assert validate_sdm_xlsx(content,spec,run,naming,release_binding=bound)['semantic_equality']=='RELEASE_LAYOUT_MATCHED'
    with pytest.raises(ValueError,match='SPECIFICATION_MISMATCH'):validate_sdm_xlsx(content,spec,run,naming)
    wb=load_workbook(BytesIO(content));old=load_workbook(BytesIO(candidate))
    assert wb.sheetnames==old.sheetnames
    assert '交付版' in wb.worksheets[0]['A3'].value
    for row in range(5,9):
        assert [c.value for c in wb.worksheets[0][row]]==[c.value for c in old.worksheets[0][row]]
    assert wb.worksheets[1]['B18'].value=='a'*64
    assert wb.worksheets[1]['B19'].value=='b'*64
    assert wb.worksheets[1]['B20'].value=='c'*64


def test_wrong_release_document_specification_rejected():
    spec,run,naming=design();bound=binding(spec);bound['specification_checksum']='0'*64
    with pytest.raises(ValueError,match='BINDING_CHANGED'):render_sdm_xlsx(spec,run,naming,release_binding=bound)


def test_release_file_exclusive_save_and_integrity(tmp_path):
    from uuid import uuid4
    ids=[uuid4() for _ in range(3)];content=b'synthetic zip bytes'
    saved=release_files.save(tmp_path,*ids,content)
    assert release_files.read(tmp_path,*ids,**saved)==content
    with pytest.raises(FileExistsError):release_files.save(tmp_path,*ids,content)
    with pytest.raises(ValueError,match='CHANGED_OR_UNAVAILABLE'):
        release_files.read(tmp_path,*ids,'0'*64,len(content))
