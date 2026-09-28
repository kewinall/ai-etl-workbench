"""JSON SDM candidate contracts; no execution or release authorization."""
from copy import deepcopy
from io import BytesIO
import json
import pytest
from openpyxl import load_workbook
from app.sdm_specification import build_sdm_candidate
from app.sdm_renderer import render_sdm_xlsx
from app.sdm_xlsx_semantics import validate_sdm_xlsx
from test_json_specification import json_design, CONTENT


@pytest.mark.parametrize('content',[CONTENT,b'\xef\xbb\xbf'+CONTENT])
def test_json_sdm_policy_mapping_and_determinism(content):
    args=json_design(content)
    before=deepcopy(args)
    candidate=build_sdm_candidate(*args)
    document=candidate['document']
    assert document['version']==5 and document['source_format']=='JSON'
    assert document['json_source']==args[0]['json_source']
    assert candidate==build_sdm_candidate(*args) and args==before
    assert not candidate['qa_passed'] and not candidate['release_ready']
    assert args[1]['input_snapshot']['source_config']['sources'][0]['upload_id'] not in json.dumps(document)
    rendered=render_sdm_xlsx(*args)
    assert rendered['content']==render_sdm_xlsx(*args)['content']
    assert rendered['evidence']['semantic_equality']=='CANDIDATE_LAYOUT_MATCHED'
    book=load_workbook(BytesIO(rendered['content']))
    rules=dict(book['規則及版本'].iter_rows(min_row=6,values_only=True))
    assert rules['JSON 根結構']=='ARRAY'
    assert 'HOP_JSON_INPUT_INCLUDE_NULLS=Y' in rules['全空物件']
    assert '不重新序列化' in rules['JSON 編碼與 BOM']
    assert rules['JSON 來源 SHA-256']==args[0]['json_source']['content_checksum']
    assert '不代表已執行或 QA 通過' in rules['讀取副本證據']
    assert '來源原名'==book['欄位對照']['E5'].value
    book.close()


@pytest.mark.parametrize('key',['content_checksum','profile_checksum','contract_checksum'])
def test_json_sdm_rejects_changed_reference(key):
    args=json_design()
    args[0]['json_source'][key]='0'*64
    with pytest.raises(ValueError,match='SDM_VALID_SPECIFICATION_REQUIRED'):
        render_sdm_xlsx(*args)


def test_json_sdm_old_source_document_rejected():
    args=json_design()
    content=render_sdm_xlsx(*args)['content']
    changed=json_design(b'\xef\xbb\xbf'+CONTENT)
    with pytest.raises(ValueError,match='SDM_XLSX_SPECIFICATION_MISMATCH'):
        validate_sdm_xlsx(content,*changed)
