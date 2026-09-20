from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy
from openpyxl import load_workbook
from app.sdm_renderer import render_sdm_xlsx
from app.sdm_xlsx_structure import PARTS
from test_etl_specification import design
import pytest
from app.platform_harness import checksum


def test_native_renderer_roundtrip_and_deterministic_bytes():
    spec, run, naming = design()
    before = deepcopy((spec, run, naming))
    result = render_sdm_xlsx(spec, run, naming)
    assert result['content'] == render_sdm_xlsx(spec, run, naming)['content']
    assert (spec, run, naming) == before
    assert not result['qa_passed'] and not result['release_ready']
    assert result['evidence']['semantic_equality'] == 'CANDIDATE_LAYOUT_MATCHED'
    with ZipFile(BytesIO(result['content'])) as archive:
        assert set(archive.namelist()) == PARTS
        assert all(b'<f>' not in archive.read(name) for name in PARTS)
    workbook = load_workbook(BytesIO(result['content']))
    assert workbook.sheetnames == ['欄位對照', '規則及版本']
    assert workbook['欄位對照']['B7'].value == 'total_amount'
    assert workbook['欄位對照']['C7'].value == 'NUMERIC(18,2)'
    assert workbook['欄位對照']['A6'].value == 1
    assert workbook['欄位對照'].freeze_panes == 'A6'
    assert workbook['規則及版本']['B12'].value == str(run['run_id'])


@pytest.mark.parametrize('source_name', ['=1+1', '+SUM(A1)', '-1', '@name'])
def test_formula_looking_source_names_are_literal(source_name):
    spec, run, naming = design()
    run['input_snapshot']['source_config']['sources'][0]['fields'][0]['name'] = source_name
    naming['contract_json']['columns'][0]['source_name'] = source_name
    naming['checksum'] = checksum(naming['contract_json']['columns'])
    spec['naming']['checksum'] = naming['checksum']
    result = render_sdm_xlsx(spec, run, naming)
    workbook = load_workbook(BytesIO(result['content']), data_only=False)
    cell = workbook['欄位對照']['E6']
    assert cell.value == source_name and cell.data_type == 's'
    assert all(cell.data_type != 'f' for sheet in workbook for row in sheet for cell in row)


def test_stale_specification_cannot_generate_excel():
    spec, run, naming = design()
    run['matches_current'] = False
    with pytest.raises(ValueError):
        render_sdm_xlsx(spec, run, naming)
