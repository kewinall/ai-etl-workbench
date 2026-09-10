from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
UPLOAD_ROOT = ROOT / "runtime-temp" / "task-uploads"


def _field_type(values: list[Any]) -> str:
    clean = [v for v in values if v not in (None, "")]
    if not clean:
        return "VARCHAR(255)"
    if all(isinstance(v, bool) or str(v).lower() in ("true", "false") for v in clean):
        return "BOOLEAN"
    try:
        if all(str(v).lstrip("+-").isdigit() for v in clean): return "BIGINT"
        if all(float(v) is not None for v in clean): return "DECIMAL(18,4)"
    except (TypeError, ValueError):
        pass
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            if all(datetime.strptime(str(v), fmt) for v in clean): return "TIMESTAMP" if " " in fmt else "DATE"
        except ValueError:
            continue
    length = min(4000, max(32, max(len(str(v)) for v in clean)))
    return f"VARCHAR({length})"


def _csv_profile(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()[:65536]
    encoding = "utf-8-sig"
    try: raw.decode(encoding)
    except UnicodeDecodeError: encoding = "big5"
    text = raw.decode(encoding, errors="replace")
    try: delimiter = csv.Sniffer().sniff(text, delimiters=",\t;|").delimiter
    except csv.Error: delimiter = ","
    with path.open("r", encoding=encoding, newline="") as stream:
        reader = csv.DictReader(stream, delimiter=delimiter)
        rows = [row for _, row in zip(range(20), reader)]
        names = reader.fieldnames or []
    return {"encoding": encoding, "delimiter": delimiter,
            "fields": [{"name": name, "type": _field_type([row.get(name) for row in rows])} for name in names],
            "sample_rows": rows[:5]}


def _excel_profile(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    worksheet = book.sheetnames[0]; sheet = book[worksheet]
    values = list(sheet.iter_rows(min_row=1, max_row=21, values_only=True))
    headers = [str(x).strip() if x is not None else f"column_{i+1}" for i, x in enumerate(values[0] if values else [])]
    rows = [dict(zip(headers, row)) for row in values[1:]]
    return {"worksheet": worksheet, "worksheets": book.sheetnames, "header_row": 1,
            "fields": [{"name": name, "type": _field_type([row.get(name) for row in rows])} for name in headers],
            "sample_rows": rows[:5]}

def _json_profile(path: Path) -> dict[str, Any]:
    text=path.read_text(encoding="utf-8-sig");parsed=json.loads(text)
    rows=parsed if isinstance(parsed,list) else [parsed]
    if not rows or any(not isinstance(row,dict) for row in rows):raise ValueError("JSON 必須是 object 或 object array")
    names=list(dict.fromkeys(key for row in rows[:20] for key in row))
    return {"encoding":"utf-8-sig","fields":[{"name":name,"type":_field_type([row.get(name) for row in rows[:20]])} for name in names],"sample_rows":rows[:5],"parser_format":"JSON"}


def save_and_profile(filename: str, content: bytes, retention_days: int = 7, max_file_mb: int = 50) -> dict[str, Any]:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in (".csv", ".xlsx", ".json"): raise ValueError("只允許上傳 CSV、XLSX 或 JSON")
    if not content: raise ValueError("上傳檔案是空的")
    if len(content) > max_file_mb * 1024 * 1024: raise ValueError(f"檔案不可超過 {max_file_mb} MB")
    upload_id = str(uuid.uuid4()); folder = UPLOAD_ROOT / upload_id; folder.mkdir(parents=True, exist_ok=False)
    safe_name = f"source{suffix}"; path = folder / safe_name; path.write_bytes(content)
    profile = _csv_profile(path) if suffix == ".csv" else (_excel_profile(path) if suffix == ".xlsx" else _json_profile(path))
    expires = datetime.now(timezone.utc) + timedelta(days=retention_days)
    return {"upload_id": upload_id, "original_name": Path(filename).name, "path": str(path),
            "source_type": "CSV" if suffix == ".csv" else ("EXCEL" if suffix == ".xlsx" else "JSON"), "size": len(content),
            "checksum": hashlib.sha256(content).hexdigest(), "expires_at": expires.isoformat(), **profile}


def cleanup_expired(retention_days: int = 7) -> int:
    UPLOAD_ROOT.mkdir(parents=True, exist_ok=True); cutoff = datetime.now().timestamp() - retention_days * 86400; removed = 0
    for folder in UPLOAD_ROOT.iterdir():
        if folder.is_dir() and folder.stat().st_mtime < cutoff:
            shutil.rmtree(folder, ignore_errors=True); removed += 1
    return removed
