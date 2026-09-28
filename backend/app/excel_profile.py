"""Explicit workbook selection bound to managed bytes; never execution approval."""
from hashlib import sha256
from io import BytesIO
import json
from zipfile import BadZipFile

from fastapi import APIRouter, HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict, Field

from .field_inference import infer_field_type


def inspect_workbook(content, worksheet=None, header_row=1):
    if type(header_row) is not int or not 1 <= header_row <= 1000:
        raise ValueError('標頭列必須是 1–1000 的整數')
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        raise ValueError('無法讀取 XLSX 工作簿，請檢查檔案格式') from None
    try:
        names = book.sheetnames
        if worksheet is None:
            return {'worksheets': names, 'worksheet': None, 'header_row': header_row,
                    'fields': [], 'sample_rows': [], 'selection_required': True}
        if worksheet not in names:
            raise ValueError('指定的工作表不存在，請重新選擇')
        sheet = book[worksheet]
        if (sheet.max_column or 0) > 512:
            raise ValueError('目前 Excel 解析最多支援 512 欄，請先整理來源')
        values = list(sheet.iter_rows(min_row=header_row, max_row=header_row + 20, values_only=True))
        headers = [str(value).strip() if value is not None else '' for value in (values[0] if values else [])]
        if not headers or any(not name for name in headers):
            raise ValueError('標頭列有空白欄名，請選擇正確標頭列或修正來源')
        if len(set(headers)) != len(headers):
            raise ValueError('標頭列有重複欄名，請修正來源，避免資料欄位被覆蓋')
        rows = [dict(zip(headers, row)) for row in values[1:]]
        return {'worksheets': names, 'worksheet': worksheet, 'header_row': header_row,
                'fields': [{'name': name, 'type': infer_field_type([row.get(name) for row in rows])} for name in headers],
                'sample_rows': rows[:5], 'sample_limit': 20, 'sample_count': len(rows),
                'selection_required': False, 'full_file_validated': False}
    finally:
        book.close()


def selected_profile(upload_id, checksum, size, worksheet, header_row, *, content=None):
    from .upload_integrity import read_verified_upload
    if content is None:
        content = read_verified_upload(upload_id, 'EXCEL', checksum, size)
    elif len(content) != size or sha256(content).hexdigest() != checksum:
        raise ValueError('UPLOAD_CONTENT_CHANGED')
    profile = inspect_workbook(content, worksheet, header_row)
    binding = {'version': 1, 'upload_id': upload_id, 'content_checksum': checksum,
               'byte_count': size, 'worksheet': worksheet, 'header_row': header_row,
               'fields': profile['fields'], 'sample_limit': 20}
    profile['excel_selection_v1'] = {**binding, 'profile_checksum': sha256(
        json.dumps(binding, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest(),
        'scope': 'INPUT_SELECTION_ONLY'}
    return profile


def verify_excel_source(source, *, content=None):
    """New inputs must reproduce their selected profile; history is not migrated."""
    binding = source.get('excel_selection_v1')
    if not isinstance(binding, dict) or not source.get('worksheet'):
        raise ValueError('Excel 來源尚未確認工作表與標頭列')
    profile = selected_profile(source.get('upload_id'), source.get('checksum'), source.get('size'),
                               source.get('worksheet'), source.get('header_row'), content=content)
    if binding != profile['excel_selection_v1'] or source.get('fields') != profile['fields']:
        raise ValueError('Excel 欄位或選擇版本已變更，請重新解析確認')
    return profile


class ExcelProfileInput(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    checksum: str = Field(pattern=r'^[0-9a-f]{64}$')
    size: int = Field(gt=0, le=50 * 1024 * 1024)
    worksheet: str = Field(min_length=1, max_length=31)
    header_row: int = Field(ge=1, le=1000)


def create_excel_profile_router():
    router = APIRouter()

    @router.post('/api/task-sources/{upload_id}/excel-profile')
    def select_excel_profile(upload_id: str, data: ExcelProfileInput):
        try:
            return selected_profile(upload_id, **data.model_dump())
        except ValueError as error:
            raise HTTPException(422, str(error)) from None

    return router
