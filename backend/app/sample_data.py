from __future__ import annotations

import csv
import json
import os
import re
import uuid
from datetime import date, timedelta
from pathlib import Path
from typing import Any

ROOT=Path(__file__).resolve().parents[2]

def _ident(value:str)->str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*",value or ""):raise ValueError(f"Unsafe identifier: {value}")
    return value

def _value(name:str,kind:str,index:int):
    t=kind.lower();n=name.lower()
    if "date" in t:return (date(2026,1,1)+timedelta(days=index)).isoformat()
    if any(x in t for x in ("int","numeric","decimal","float","double","number")):
        if n.endswith("_id"):return 1000+(index%3)+1
        if "price" in n or "amount" in n:return round(100+(index*17.5),2)
        return index+1
    if "bool" in t:return index%2==0
    if "region" in n:return ("North","South","West")[index%3]
    if "category" in n:return ("Hardware","Software","Service")[index%3]
    return f"{name}_{index+1}"

def _rows(fields:list[dict[str,Any]],count=10):
    return [{f["name"]:_value(f["name"],f.get("type") or "varchar",i) for f in fields} for i in range(count)]

def _vertica_type(value:str)->str:
    value=(value or "VARCHAR(255)").upper()
    allowed=("INT","INTEGER","BIGINT","SMALLINT","NUMERIC","DECIMAL","FLOAT","DOUBLE","DATE","TIMESTAMP","BOOLEAN","VARCHAR","CHAR")
    if not value.startswith(allowed):raise ValueError(f"Unsupported sample field type: {value}")
    if not re.fullmatch(r"[A-Z]+(?:\(\d+(?:,\d+)?\))?",value):raise ValueError(f"Unsafe sample field type: {value}")
    return value

def _ensure_schema(cur,schema:str):
    if schema.lower()!='public':cur.execute(f'CREATE SCHEMA IF NOT EXISTS "{schema}"')

def materialize_samples(task_id:str,source_config:dict[str,Any])->dict[str,Any]:
    sources=source_config.get("sources") or []
    generated=[];runtime=ROOT/'runtime-temp'/'generated-sources'/task_id;runtime.mkdir(parents=True,exist_ok=True)
    for i,source in enumerate(sources):
        if source.get("has_actual_data",source_config.get("has_actual_data",True)) is not False:
            generated.append(source)
            continue
        fields=source.get("fields") or []
        if not fields or any(not f.get("name") or not f.get("type") for f in fields):raise ValueError(f"來源 {i+1} 必須填寫所有欄位名稱與型別")
        for f in fields:_ident(f["name"])
        rows=_rows(fields);kind=source.get("type")
        if kind=="VERTICA":
            import vertica_python
            schema=_ident(source.get("schema") or "public");table=_ident(source.get("object") or f"sample_{task_id.lower().replace('-','_')}_{i+1}")
            cfg={'host':os.getenv('VERTICA_HOST','127.0.0.1'),'port':int(os.getenv('VERTICA_PORT','5433')),'user':os.getenv('VERTICA_USER','dbadmin'),'password':os.getenv('VERTICA_PASSWORD',''),'database':os.getenv('VERTICA_DATABASE','VMart'),'tlsmode':os.getenv('VERTICA_TLSMODE','disable')}
            definitions=','.join(f'"{f["name"]}" {_vertica_type(f["type"])}' for f in fields)
            with vertica_python.connect(**cfg) as conn:
                cur=conn.cursor();_ensure_schema(cur,schema);cur.execute(f'CREATE TABLE IF NOT EXISTS "{schema}"."{table}" ({definitions})')
                cur.execute(f'DELETE FROM "{schema}"."{table}"')
                placeholders=','.join(['%s']*len(fields));insert_sql=f'INSERT INTO "{schema}"."{table}" VALUES ({placeholders})'
                for row in rows:cur.execute(insert_sql,[row[f['name']] for f in fields])
                conn.commit()
            source={**source,"schema":schema,"object":table}
        elif kind=="CSV":
            path=runtime/f'source_{i+1}.csv'
            with path.open('w',encoding='utf-8-sig',newline='') as stream:
                writer=csv.DictWriter(stream,fieldnames=[f['name'] for f in fields]);writer.writeheader();writer.writerows(rows)
            source={**source,"path":str(path),"file_path":str(path),"encoding":"utf-8-sig","delimiter":","}
        elif kind=="EXCEL":
            from openpyxl import Workbook
            path=runtime/f'source_{i+1}.xlsx';book=Workbook();sheet=book.active;sheet.title=source.get('worksheet') or 'Data';sheet.append([f['name'] for f in fields])
            for row in rows:sheet.append([row[f['name']] for f in fields])
            book.save(path);source={**source,"path":str(path),"worksheet":sheet.title,"header_row":1,"formula_values":True}
        elif kind=="JSON":
            path=runtime/f'source_{i+1}.json'
            path.write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
            source={**source,"path":str(path),"file_path":str(path),"parser_format":"JSON"}
        else:raise ValueError(f"Unsupported generated source type: {kind}")
        generated.append(source)
    return {**source_config,"sources":generated,"generated_sample":True,"generated_row_count":10}

def rebuild_managed_sample(schema:str,table:str,fields:list[dict[str,Any]],sample_rows:list[dict[str,Any]]|None=None,allow_replace:bool=False)->dict[str,Any]:
    """Rebuild a platform-approved sample table. Caller must authorize ownership first."""
    schema=_ident(schema);table=_ident(table)
    if schema.lower()!='ai_sample':raise ValueError('Only ai_sample schema supports drop/recreate')
    if not fields or any(not item.get('name') or not item.get('type') for item in fields):raise ValueError('範例表必須提供完整欄位名稱與型別')
    for item in fields:_ident(item['name']);_vertica_type(item['type'])
    rows=sample_rows or _rows(fields);definitions=','.join(f'"{_ident(item["name"])}" {_vertica_type(item["type"])}' for item in fields)
    ddl=f'CREATE TABLE "{schema}"."{table}" ({definitions})'
    cfg={'host':os.getenv('VERTICA_HOST','127.0.0.1'),'port':int(os.getenv('VERTICA_PORT','5433')),'user':os.getenv('VERTICA_USER','dbadmin'),'password':os.getenv('VERTICA_PASSWORD',''),'database':os.getenv('VERTICA_DATABASE','VMart'),'tlsmode':os.getenv('VERTICA_TLSMODE','disable')}
    with __import__('vertica_python').connect(**cfg) as conn:
        cur=conn.cursor();_ensure_schema(cur,schema);cur.execute("SELECT 1 FROM tables WHERE table_schema=%s AND table_name=%s",(schema,table));exists=bool(cur.fetchone())
        if exists and not allow_replace:raise ValueError('Existing table is not platform-managed; refusing to drop it')
        if exists:cur.execute(f'DROP TABLE "{schema}"."{table}"')
        cur.execute(ddl)
        placeholders=','.join(['%s']*len(fields));insert_sql=f'INSERT INTO "{schema}"."{table}" VALUES ({placeholders})'
        for row in rows:cur.execute(insert_sql,[row.get(item['name']) for item in fields])
        conn.commit()
    return {'schema':schema,'table':table,'fields':fields,'sample_rows':rows,'ddl':ddl,'row_count':len(rows)}

def ensure_sample_target(target:dict[str,Any],fields:list[dict[str,Any]]):
    import vertica_python
    schema=_ident(target.get('schema') or 'public');table=_ident(target.get('table') or 'generated_result')
    definitions=','.join(f'"{_ident(f["name"])}" {_vertica_type(f["type"])}' for f in fields)
    cfg={'host':os.getenv('VERTICA_HOST','127.0.0.1'),'port':int(os.getenv('VERTICA_PORT','5433')),'user':os.getenv('VERTICA_USER','dbadmin'),'password':os.getenv('VERTICA_PASSWORD',''),'database':os.getenv('VERTICA_DATABASE','VMart'),'tlsmode':os.getenv('VERTICA_TLSMODE','disable')}
    with vertica_python.connect(**cfg) as conn:
        cur=conn.cursor();_ensure_schema(cur,schema);cur.execute(f'CREATE TABLE IF NOT EXISTS "{schema}"."{table}" ({definitions})')

def ensure_query_target(target:dict[str,Any],query:str):
    import vertica_python
    schema=_ident(target.get('schema') or 'public');table=_ident(target.get('table') or 'generated_result')
    statement=(query or '').strip().rstrip(';').strip()
    if not re.match(r'(?is)^(select|with)\b',statement) or ';' in statement:
        raise ValueError('Target bootstrap query must be one read-only SELECT/WITH statement')
    cfg={'host':os.getenv('VERTICA_HOST','127.0.0.1'),'port':int(os.getenv('VERTICA_PORT','5433')),'user':os.getenv('VERTICA_USER','dbadmin'),'password':os.getenv('VERTICA_PASSWORD',''),'database':os.getenv('VERTICA_DATABASE','VMart'),'tlsmode':os.getenv('VERTICA_TLSMODE','disable')}
    with vertica_python.connect(**cfg) as conn:
        cur=conn.cursor();_ensure_schema(cur,schema)
        cur.execute(f'CREATE TABLE IF NOT EXISTS "{schema}"."{table}" AS SELECT * FROM ({statement}) AS target_bootstrap WHERE 1=0')

def materialize_stage_server_sample(task_id:str,target:dict[str,Any],source:dict[str,Any])->dict[str,Any]:
    """Create a Task-scoped CSV/JSON on the Vertica server for External/Flex modes."""
    mode=str(target.get("stage_mode") or "NORMAL").upper()
    if mode not in ("EXTERNAL","FLEX"):return target
    fields=source.get("fields") or []
    if not fields or any(not f.get("name") or not f.get("type") for f in fields):raise ValueError("Server 範例檔需要完整欄位名稱與型別")
    for item in fields:_ident(item["name"]);_vertica_type(item["type"])
    base=str(target.get("data_directory") or "").rstrip("/")
    if not base.startswith("/") or any(x in base for x in ("'",";","*","?","\n","\r")):raise ValueError("Data Folder 必須是安全的 Vertica Server 絕對資料夾")
    suffix=uuid.uuid4().hex[:8];directory=f"{base}/{task_id.lower()}_{suffix}";extension="json" if mode=="FLEX" and str(target.get("format") or "JSON").upper()=="JSON" else "csv";filename=f"source_1.{extension}"
    schema="public";table=_ident(f"stage_sample_{task_id.lower().replace('-','_')}_{suffix}");definitions=','.join(f'"{_ident(f["name"])}" {_vertica_type(f["type"])}' for f in fields)
    path=Path(str(source.get("path") or source.get("file_path") or ""));rows=_rows(fields)
    if source.get("has_actual_data",True) is not False:
        if not path.is_file():raise ValueError("上傳檔案不存在或已過期")
        if path.suffix.lower()==".json":
            import json as json_module
            parsed=json_module.loads(path.read_text(encoding="utf-8-sig"));rows=parsed if isinstance(parsed,list) else [parsed]
        elif path.suffix.lower()==".csv":
            with path.open("r",encoding=source.get("encoding") or "utf-8-sig",newline="") as stream:rows=list(csv.DictReader(stream,delimiter=source.get("delimiter") or ",",quotechar=source.get("quote") or '"'))
        else:raise ValueError("External／Flex 上傳只支援 CSV 或 JSON")
    cfg={'host':os.getenv('VERTICA_HOST','127.0.0.1'),'port':int(os.getenv('VERTICA_PORT','5433')),'user':os.getenv('VERTICA_USER','dbadmin'),'password':os.getenv('VERTICA_PASSWORD',''),'database':os.getenv('VERTICA_DATABASE','VMart'),'tlsmode':os.getenv('VERTICA_TLSMODE','disable')}
    with __import__('vertica_python').connect(**cfg) as conn:
        target_schema=_ident(target.get('schema') or 'public')
        cur=conn.cursor();_ensure_schema(cur,target_schema);cur.execute(f'CREATE TABLE "{schema}"."{table}" ({definitions})');placeholders=','.join(['%s']*len(fields));insert_sql=f'INSERT INTO "{schema}"."{table}" VALUES ({placeholders})'
        for row in rows:cur.execute(insert_sql,[row[f['name']] for f in fields])
        conn.commit()
        select_list=','.join(f'"{_ident(f["name"])}"' for f in fields)
        if extension=="json":sql=f"EXPORT TO JSON(directory='{directory}', filename='{filename}') AS SELECT {select_list} FROM \"{schema}\".\"{table}\""
        else:sql=f"EXPORT TO DELIMITED(directory='{directory}', filename='{filename}', delimiter=',', enclosedBy='\"', addHeader='true') AS SELECT {select_list} FROM \"{schema}\".\"{table}\""
        try:cur.execute(sql);cur.fetchone()
        finally:cur.execute(f'DROP TABLE IF EXISTS "{schema}"."{table}"');conn.commit()
    return {**target,"data_path":f"{directory}/*.{extension}","server_directory":directory,"server_file":filename}
