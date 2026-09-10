from __future__ import annotations

import csv
import hashlib
import json
import os
import re
from pathlib import Path
from typing import Any
from .platform_harness import inferred_vertica_type


def normalized_sources(task: dict[str, Any]) -> list[dict[str, Any]]:
    config = task.get("source_config") or {}
    configured = config.get("sources")
    if isinstance(configured, list) and configured:
        return [{**source, "type": str(source.get("type") or "").upper()} for source in configured]
    source_type = str(task.get("source") or "CSV").upper()
    if source_type == "CSV":
        return [{"id": "source_1", "alias": "src", "type": "CSV", **config,
                 "path": config.get("path") or config.get("file_path")}]
    if source_type in ("EXCEL","JSON"):
        return [{"id": "source_1", "alias": "src", "type": "EXCEL", **config,
                 "path": config.get("path") or config.get("file_path"),"type":source_type}]
    if source_type in ("VERTICA", "VERTICA_TABLE"):
        return [{"id": "source_1", "alias": "src", "type": "VERTICA", **config,
                 "object": config.get("object") or config.get("table")}]
    return configured or []


def _checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _masked(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    return text[:2] + "***" if text else text


def _profile_csv(source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    path = Path(str(source.get("path") or source.get("file_path") or ""))
    if not path.is_file():
        raise ValueError(f"CSV file not found: {path}")
    encoding = source.get("encoding") or "utf-8-sig"
    delimiter = source.get("delimiter")
    with path.open("r", encoding=encoding, newline="") as stream:
        head = stream.read(8192)
        stream.seek(0)
        if not delimiter:
            try: delimiter = csv.Sniffer().sniff(head, delimiters=",\t;|").delimiter
            except csv.Error: delimiter = None
        if not delimiter:
            return ({**source, "checksum": _checksum(path), "fields": []}, [])
        reader = csv.DictReader(stream, delimiter=delimiter, quotechar=source.get("quote") or '"')
        rows = []
        malformed = 0
        for index, row in enumerate(reader):
            if None in row: malformed += 1
            if index < 10: rows.append(dict(row))
        fields = [{"name": name, "type": inferred_vertica_type([row.get(name) for row in rows])} for name in (reader.fieldnames or [])]
    return ({**source, "path": str(path), "encoding": encoding, "delimiter": delimiter,
             "checksum": _checksum(path), "fields": fields, "row_count": index + 1 if 'index' in locals() else 0,
             "malformed_rows": malformed, "masked_examples": [{k: _masked(v) for k, v in r.items()} for r in rows[:3]]}, rows)


def _profile_excel(source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel profiling requires openpyxl") from exc
    path = Path(str(source.get("path") or source.get("file_path") or ""))
    if not path.is_file(): raise ValueError(f"Excel file not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=bool(source.get("formula_values", True)))
    worksheet = source.get("worksheet")
    if worksheet not in workbook.sheetnames:
        if worksheet: raise ValueError(f"Worksheet not found: {worksheet}")
        worksheet = workbook.sheetnames[0] if len(workbook.sheetnames) == 1 else None
    if not worksheet:
        return ({**source, "path": str(path), "worksheets": workbook.sheetnames,
                 "checksum": _checksum(path), "fields": []}, [])
    sheet = workbook[worksheet]; header_row = int(source.get("header_row") or 1)
    values = list(sheet.iter_rows(min_row=header_row, max_row=header_row + 10, values_only=True))
    headers = [str(value).strip() if value is not None else f"column_{i+1}" for i, value in enumerate(values[0] if values else [])]
    rows = [dict(zip(headers, row)) for row in values[1:]]
    return ({**source, "path": str(path), "worksheet": worksheet, "worksheets": workbook.sheetnames,
             "header_row": header_row, "checksum": _checksum(path),
             "fields": [{"name": h, "type": inferred_vertica_type([row.get(h) for row in rows])} for h in headers],
             "masked_examples": [{k: _masked(v) for k, v in r.items()} for r in rows[:3]]}, rows)

def _profile_json(source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    path=Path(str(source.get("path") or source.get("file_path") or ""))
    if not path.is_file():raise ValueError(f"JSON file not found: {path}")
    parsed=json.loads(path.read_text(encoding=source.get("encoding") or "utf-8-sig"));rows=parsed if isinstance(parsed,list) else [parsed]
    if not rows or any(not isinstance(row,dict) for row in rows):raise ValueError("JSON must contain objects")
    names=list(dict.fromkeys(key for row in rows[:20] for key in row))
    return ({**source,"path":str(path),"checksum":_checksum(path),"fields":[{"name":name,"type":inferred_vertica_type([row.get(name) for row in rows[:20]])} for name in names],"masked_examples":[{k:_masked(v) for k,v in row.items()} for row in rows[:3]]},rows[:10])


def _vertica_connection():
    import vertica_python
    return vertica_python.connect(host=os.getenv("VERTICA_HOST", "127.0.0.1"), port=int(os.getenv("VERTICA_PORT", "5433")),
                                  user=os.getenv("VERTICA_USER", "dbadmin"), password=os.getenv("VERTICA_PASSWORD", ""),
                                  database=os.getenv("VERTICA_DATABASE", "VMart"), tlsmode=os.getenv("VERTICA_TLSMODE", "disable"))


def _profile_vertica(source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    schema, object_name = source.get("schema"), source.get("object") or source.get("table")
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", schema or "") or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", object_name or ""):
        raise ValueError("Unsafe Vertica schema or object identifier")
    with _vertica_connection() as conn:
        cur=conn.cursor();cur.execute("SELECT column_name,data_type FROM columns WHERE table_schema=%s AND table_name=%s ORDER BY ordinal_position",(schema,object_name));metadata=cur.fetchall()
        if not metadata:raise ValueError(f"Vertica object not found or has no columns: {schema}.{object_name}")
        cur.execute(f'SELECT * FROM "{schema}"."{object_name}" LIMIT 3');rows=cur.fetchall();names=[x[0] for x in cur.description];sample=[dict(zip(names,row)) for row in rows]
        try:
            cur.execute("SELECT COUNT(*) FROM v_catalog.projections WHERE anchor_table_name=%s",(object_name,));projection_count=cur.fetchone()[0]
        except Exception:
            projection_count=None
    return ({**source,"object":object_name,"fields":[{"name":n,"type":t} for n,t in metadata],"projection_count":projection_count,
             "masked_examples":[{k:_masked(v) for k,v in row.items()} for row in sample]},sample)


def profile_task(task: dict[str, Any]) -> dict[str, Any]:
    sources = normalized_sources(task); profiles = []; samples = []; issues = []
    if not sources:
        issues.append({"severity": "BLOCKING", "field": "sources", "reason": "至少需要一個來源",
                       "candidates": [], "suggested_default": None})
    for source in sources:
        kind = source.get("type")
        if kind == "CSV": profile, rows = _profile_csv(source)
        elif kind == "EXCEL": profile, rows = _profile_excel(source)
        elif kind == "JSON": profile, rows = _profile_json(source)
        elif kind == "VERTICA":
            required = [key for key in ("connection", "schema", "object") if not source.get(key)]
            profile, rows = _profile_vertica(source) if not required else ({**source, "fields": []}, [])
            if required:
                issues.append({"severity": "BLOCKING", "field": f"sources.{required[0]}",
                               "reason": "Vertica 來源缺少連線、schema 或 object", "candidates": [],
                               "suggested_default": None})
        else: raise ValueError(f"Unsupported source type: {kind}")
        if kind == "CSV" and not profile.get("delimiter"):
            issues.append({"severity": "BLOCKING", "field": "sources.delimiter",
                           "reason": "無法可靠偵測 CSV delimiter", "candidates": [",", "\\t", ";", "|"],
                           "suggested_default": ","})
        if kind == "CSV" and profile.get("malformed_rows"):
            issues.append({"severity": "BLOCKING", "field": "sources.column_count",
                           "reason": "CSV 存在欄數不一致資料列", "candidates": [], "suggested_default": None})
        if kind == "EXCEL" and not profile.get("worksheet"):
            issues.append({"severity": "BLOCKING", "field": "sources.worksheet",
                           "reason": "Excel 有多個 worksheet，必須指定工作表", "candidates": profile.get("worksheets", []),
                           "suggested_default": profile.get("worksheets", [None])[0]})
        profiles.append(profile); samples.extend(rows[:10])
    fingerprint = hashlib.sha256("|".join(str(s.get("checksum") or f"{s.get('schema')}.{s.get('object')}") for s in profiles).encode()).hexdigest()
    configured_output=(task.get("source_config") or {}).get("output_fields") or []
    columns=[x.get("name") if isinstance(x,dict) else x for x in configured_output] or [f["name"] for s in profiles for f in s.get("fields", [])]
    return {"sources": profiles, "samples": samples[:10], "issues": issues, "fingerprint": fingerprint,
            "columns": columns}
