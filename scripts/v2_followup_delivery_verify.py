"""Download and verify the three formal bundles without rerunning ETL."""
import json, hashlib, urllib.request, io, zipfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
state=json.loads((ROOT/'runtime-temp/v2-followup-live-state.json').read_text(encoding='utf-8'))
folder=ROOT/'runtime-temp/v2-followup-formal-downloads';folder.mkdir(exist_ok=True)
records=[]
def get(url):
    with urllib.request.urlopen('http://127.0.0.1:5197/api'+url,timeout=30) as r:return json.load(r)
for key,c in state['cases'].items():
    base='/tasks/'+c['task_id']+'/runs/'+c['run_id']
    d=get(base+'/release')
    assert d['status']=='RELEASE_READY' and d['release_ready'] and d['portability']['status']=='PASS',key+' NOT_RELEASE_READY'
    qa=get(base+'/qa-review');assert qa['invocation']['review']['status']=='PASS' and get(base+'/qa-approval')['status']=='APPROVED_CURRENT'
    models={}
    for endpoint,role in [('sa-invocation','SA'),('developer','Developer'),('qa-review','QA')]:
        i=get(base+'/'+endpoint)['invocation'];assert i['provider']=='LOCAL_COPILOT' and i['status']=='VALIDATED_NOT_APPROVED'
        assert i['usage']['cli_sessions']==1 and i['usage']['automatic_retries']==0 and i['usage']['tool_execution_count']==0
        models[role]={k:i[k] for k in ['invocation_id','provider','model','duration_ms','usage']}
    assert get(base+'/sa-approval')['status']=='APPROVED_CURRENT'
    release=d['release']
    with urllib.request.urlopen('http://127.0.0.1:5197/api'+base+'/release/'+release['release_id']+'/download') as r:raw=r.read()
    digest=hashlib.sha256(raw).hexdigest();assert digest==release['checksum'] and len(raw)==release['file_size']
    archive=zipfile.ZipFile(io.BytesIO(raw));assert archive.testzip() is None
    manifest_bytes=archive.read('release-manifest.json');manifest=json.loads(manifest_bytes)
    assert manifest['release_ready'] and manifest['qa_passed'] and manifest['run_id']==c['run_id']
    assert manifest['portability']=='ISOLATED_HOP_RESULT_VERIFIED'
    bound=manifest['approval_binding'];assert bound['portability_checksum']==d['portability']['checksum'] and bound['candidate_checksum']==d['candidate']['checksum']
    expected={'hop/pipeline.hpl','hop/workflow.hwf','vertica-ddl.sql','SDM.xlsx','parameters.example'}
    assert {a['name'] for a in manifest['artifacts']}==expected and set(archive.namelist())==expected|{'release-manifest.json'}
    for a in manifest['artifacts']:
        data=archive.read(a['name']);assert hashlib.sha256(data).hexdigest()==a['checksum'] and len(data)==a['size_bytes']
    workbook=load_workbook(io.BytesIO(archive.read('SDM.xlsx')),read_only=True)
    assert workbook.sheetnames; sheets=workbook.sheetnames;workbook.close()
    (folder/(key+'-release.zip')).write_bytes(raw);(folder/(key+'-SDM.xlsx')).write_bytes(archive.read('SDM.xlsx'))
    workflow=get(base+'/workflow-status');assert workflow['release_status']=='RELEASE_READY' and len(workflow['completed'])==7
    sdms=get(base+'/sdm-candidates');assert sdms['items']
    record={'key':key,**{field:c[field] for field in ('task_id','run_id','target','specification_id')},'models':models,'qa':'PASS','release':release,'portability':d['portability'],'manifest_sha256':hashlib.sha256(manifest_bytes).hexdigest(),'all_five_artifact_hashes_verified':True,'sdm_sheets':sheets,'workflow_completed_stages':7,'download_url':'http://127.0.0.1:5197/api'+base+'/release/'+release['release_id']+'/download'}
    records.append(record)
proof={'scope':'Three true native cases, official API downloads and all manifest members; no ETL replay','cases':records,'all_three_formally_delivered':True}
(ROOT/'docs/verification/v2-followup-delivery-2026-10-11.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps([{'case':r['key'],'release_id':r['release']['release_id'],'bytes':r['release']['file_size'],'sha256':r['release']['checksum']} for r in records]))
