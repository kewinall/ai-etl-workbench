"""GET-only final proof for the authorized failed-parent/corrected-child case."""
import hashlib,io,json,urllib.request,zipfile
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
s=json.loads((root/'runtime-temp/v2-revision-live-state.json').read_text(encoding='utf-8-sig'))
assert s['project_id']=='346e9419-4b0e-4ba7-ad29-1756ca7da833' and s['target']=='ui_accept_revision_20261011_v2'
def get(path,raw=False):
 with urllib.request.urlopen('http://127.0.0.1:5197/api'+path,timeout=30) as response:
  return response.read() if raw else json.load(response)
base='/tasks/'+s['task_id']+'/runs/'+s['run_id']
d=get(base+'/release');assert d['release_ready'] and d['status']=='RELEASE_READY'
assert d['portability']['status']=='PASS' and get(base+'/qa-approval')['status']=='APPROVED_CURRENT'
models={}
for endpoint,role in [('sa-invocation','SA'),('developer','Developer'),('qa-review','QA')]:
 i=get(base+'/'+endpoint)['invocation']
 assert i['provider']=='LOCAL_COPILOT' and i['status']=='VALIDATED_NOT_APPROVED'
 assert i['usage']['cli_sessions']==1 and i['usage']['automatic_retries']==0
 models[role]={k:i[k] for k in ('invocation_id','provider','model','duration_ms','usage')}
assert get(base+'/qa-review')['invocation']['review']['status']=='PASS'
release=d['release'];raw=get(base+'/release/'+release['release_id']+'/download',True)
assert len(raw)==release['file_size'] and hashlib.sha256(raw).hexdigest()==release['checksum']
with zipfile.ZipFile(io.BytesIO(raw)) as archive:
 assert archive.testzip() is None
 manifest_raw=archive.read('release-manifest.json');manifest=json.loads(manifest_raw)
 assert manifest['run_id']==s['run_id'] and manifest['release_ready'] and manifest['qa_passed']
 assert manifest['portability']=='ISOLATED_HOP_RESULT_VERIFIED'
 assert manifest['approval_binding']['portability_checksum']==d['portability']['checksum']
 assert manifest['approval_binding']['candidate_checksum']==d['candidate']['checksum']
 members={'hop/pipeline.hpl','hop/workflow.hwf','vertica-ddl.sql','SDM.xlsx','parameters.example'}
 assert set(archive.namelist())==members|{'release-manifest.json'}
 assert {a['name'] for a in manifest['artifacts']}==members
 for a in manifest['artifacts']:
  data=archive.read(a['name']);assert hashlib.sha256(data).hexdigest()==a['checksum'] and len(data)==a['size_bytes']
 workbook=load_workbook(io.BytesIO(archive.read('SDM.xlsx')),read_only=True);sheets=workbook.sheetnames;workbook.close();assert sheets
parent=get('/tasks/'+s['task_id']+'/runs/'+s['failed_execution']['run_id'])
child=get(base);assert child['parent_run_id']==parent['run_id']
assert parent['outcome_code']=='HOP_EXECUTION_FAILED' and sum(e['event_type']=='WRITE_STARTED' for e in parent['events'])==1
assert sum(e['event_type']=='WRITE_STARTED' for e in child['events'])==1
assert get('/tasks/'+s['task_id']+'/runs/'+parent['run_id']+'/reconciliation')['status']=='CLOSED_WITHOUT_RETRY'
workflow=get(base+'/workflow-status');assert len(workflow['completed'])==7
assert workflow['operation_qualifications']['release_download']['enabled'] is True
assert all(value.get('enabled',False) is False for key,value in workflow['operation_qualifications'].items() if key!='release_download')
proof=dict(scope='5197 authorized new failure/revision case; website operations at 768px; original Hop once per Run; GET-only final proof',task_id=s['task_id'],run_id=s['run_id'],target=s['target'],parent_run_id=parent['run_id'],expected_rows=s['expected_rows'],models=models,release=release,manifest_sha256=hashlib.sha256(manifest_raw).hexdigest(),all_five_artifact_hashes_verified=True,sdm_sheets=sheets,workflow_completed_stages=7,actions=s['actions'],retained_runs=s['retained_runs'],failure=json.loads((root/'runtime-temp/v2-revision-failure-readback.json').read_text()),native=json.loads((root/'runtime-temp/v2-revision-native-readback.json').read_text()),portable=json.loads((root/'runtime-temp/v2-revision-portable-readback.json').read_text()),preserved_releases=json.loads((root/'runtime-temp/v2-revision-release-regression/after.json').read_text()),status='PASS')
proof['cases']=[dict(key='failure_revision',task_id=s['task_id'],run_id=s['run_id'],release=release)]
(root/'docs/verification/v2-revision-delivery-2026-10-11.json').write_text(json.dumps(proof,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print(json.dumps(dict(task_id=s['task_id'],run_id=s['run_id'],release_id=release['release_id'],bytes=len(raw),sha256=release['checksum'],status='PASS')))
